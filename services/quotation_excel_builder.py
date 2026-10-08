import os
from datetime import datetime
from openpyxl import load_workbook

TEMPLATE_PATH = "templates/Quo PWS-01.xlsx"
OUTPUT_DIR = "output_files"

def render_quotation_excel(quotation_json: dict):
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb[quotation_json["template_binding"].get("sheet_name", "Quotation")]

    header = quotation_json.get("header_fields", {})
    summary = quotation_json.get("summary", {})
    line_items = quotation_json.get("line_items", [])
    assumptions = quotation_json.get("assumptions_section", [])
    exclusions = quotation_json.get("exclusions_section", [])

    # ví dụ mapping cell
    ws["B4"] = header.get("customer_name", "")
    ws["B5"] = header.get("project_name", "")
    ws["B6"] = header.get("subject", "")

    start_row = int(quotation_json["template_binding"].get("start_row_line_items", 15))

    for i, item in enumerate(line_items):
        r = start_row + i
        ws[f"A{r}"] = item.get("line_no")
        ws[f"B{r}"] = item.get("item_name")
        ws[f"C{r}"] = item.get("description")
        ws[f"D{r}"] = item.get("unit")
        ws[f"E{r}"] = item.get("quantity")
        ws[f"F{r}"] = item.get("unit_price")
        ws[f"G{r}"] = item.get("amount")

    # summary ví dụ
    ws["G40"] = summary.get("subtotal", 0)
    ws["G41"] = summary.get("discount_amount", 0)
    ws["G42"] = summary.get("tax_amount", 0)
    ws["G43"] = summary.get("grand_total", 0)

    # assumptions
    base_row = 46
    for i, txt in enumerate(assumptions):
        ws[f"B{base_row+i}"] = f"- {txt}"

    # exclusions
    base_row_ex = 52
    for i, txt in enumerate(exclusions):
        ws[f"B{base_row_ex+i}"] = f"- {txt}"

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"quotation_{ts}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, filename)
    wb.save(out_path)
    return out_path, filename